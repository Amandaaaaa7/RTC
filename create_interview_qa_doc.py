from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = Path("熊晨涵科研面试模拟问答_博士资格版.docx")


def set_run_font(run, name="Microsoft YaHei", size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def set_para_format(paragraph, before=0, after=5, line=1.35):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def shade_paragraph(paragraph, fill="EEF3F8"):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run_font(run, size=9, color=(100, 100, 100))


def add_label_paragraph(doc, label, content, label_color=(35, 78, 121)):
    p = doc.add_paragraph()
    set_para_format(p, after=6, line=1.38)
    r = p.add_run(label)
    set_run_font(r, size=10.5, bold=True, color=label_color)
    r = p.add_run(content)
    set_run_font(r, size=10.5)
    return p


def add_question(doc, index, question, answer, tips=None):
    p = doc.add_paragraph()
    set_para_format(p, before=8, after=4, line=1.28)
    shade_paragraph(p)
    r = p.add_run(f"{index}. {question}")
    set_run_font(r, size=11, bold=True, color=(31, 73, 125))
    p = doc.add_paragraph()
    set_para_format(p, after=3, line=1.42)
    r = p.add_run("参考回答：")
    set_run_font(r, size=10.5, bold=True, color=(35, 78, 121))
    r = p.add_run(answer)
    set_run_font(r, size=10.5)
    if tips:
        p = doc.add_paragraph()
        set_para_format(p, after=6, line=1.28)
        r = p.add_run("回答提示：")
        set_run_font(r, size=9.5, bold=True, color=(112, 48, 160))
        r = p.add_run(tips)
        set_run_font(r, size=9.5, color=(89, 89, 89))


def add_english_question(doc, index, question, answer_en, answer_cn, keywords):
    p = doc.add_paragraph()
    set_para_format(p, before=8, after=4, line=1.28)
    shade_paragraph(p, "EAF2F8")
    r = p.add_run(f"{index}. {question}")
    set_run_font(r, name="Aptos", size=11, bold=True, color=(31, 73, 125))
    add_label_paragraph(doc, "English answer: ", answer_en)
    add_label_paragraph(doc, "中文翻译：", answer_cn)
    add_label_paragraph(doc, "关键词：", keywords, label_color=(112, 48, 160))


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    set_para_format(p, before=14 if level == 1 else 10, after=6, line=1.18)
    r = p.add_run(text)
    set_run_font(r, size=15 if level == 1 else 12.5, bold=True, color=(0, 0, 0))
    return p


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.65)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    for style_name in ["Title", "Heading 1", "Heading 2"]:
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.color.rgb = RGBColor(0, 0, 0)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para_format(title, before=0, after=6, line=1.0)
    r = title.add_run("科研面试模拟问答")
    set_run_font(r, size=21, bold=True, color=(0, 0, 0))
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para_format(subtitle, after=16, line=1.1)
    r = subtitle.add_run("熊晨涵  自动化专业本科生  面向科研经历与技术细节")
    set_run_font(r, size=10.5, color=(89, 89, 89))

    add_heading(doc, "使用说明")
    add_label_paragraph(doc, "适用场景：", "推免复试、科研夏令营、导师组面试以及技术类项目答辩。题目按照老师常见的追问路径编排，先问项目整体，再问方法、验证、个人贡献与局限。")
    add_label_paragraph(doc, "练习方式：", "先用自己的话回答，再对照参考答案补足逻辑。每题建议控制在 40 至 90 秒，遇到老师追问时再展开细节。")
    add_label_paragraph(doc, "重要提醒：", "简历未写明具体模型版本、帧率、延迟、准确率或数据规模。文中的回答避免虚构这些信息；带有“按真实情况补充”的位置，请在面试前换成你实际使用过的内容。")
    add_label_paragraph(doc, "通用结构：", "背景与目标 -> 核心问题 -> 你的方法与贡献 -> 效果、验证和反思。技术问题先给结论，再说明依据，避免只罗列术语。")

    add_heading(doc, "一 玩偶机器人多模态交互系统")
    robot_questions = [
        ("请你完整介绍一下玩偶机器人交互系统，你在其中主要做了什么？",
         "这个项目的目标，是让计算能力有限的玩偶机器人能够在本地完成简单语音响应，并根据人的位置变化调整注视方向。难点在于设备是树莓派 Zero 2 W，计算资源有限，不能直接照搬电脑端方案。我参与的重点是把轻量级的语音关键词检测和量化后的模型部署到设备端，并配合视觉坐标映射，让视觉、语音和眼神、声音反馈形成一个持续运行的交互闭环。项目最终实现了不依赖网络的本地实时交互。",
         "把“我负责什么”说具体：部署、流程衔接、坐标映射或调试中的哪几项是你亲手完成的。"),
        ("为什么要强调端侧部署，而不是把语音和视觉任务放到云端？",
         "这个选择首先来自使用场景。玩偶机器人需要及时回应，网络往返会带来不确定的延迟，也会让设备在断网时失去核心交互能力。其次，小型陪伴设备通常更适合把简单、频繁、对时延敏感的任务放在本地。因此我们把关键词检测等功能部署在树莓派上，让机器人能在本地完成触发和反馈。这样做的代价是模型和计算流程都要更轻量，需要在功能效果和资源占用之间做取舍。",
         "可补充你实际观察到的网络问题或现场使用需求，但不要编造具体延迟数值。"),
        ("你提到使用了关键词检测和 INT8 量化，它们分别解决什么问题？",
         "关键词检测解决的是“快速判断用户是否发出了需要回应的简单指令”。相比开放式语音识别，它聚焦于有限的触发词，因此更适合资源有限的设备。INT8 量化解决的是模型部署时的存储和计算开销问题。它会用更紧凑的数值表示模型参数和计算结果，从而减少内存占用，并让设备端推理更轻。我的理解是，量化不是为了追求复杂功能，而是为了让已有功能能稳定地跑在目标硬件上。",
         "如果老师问“量化会不会损失效果”，可答：会有精度和鲁棒性的权衡，需要在目标词识别表现和实际响应速度之间做测试。"),
        ("视觉注视是怎么从摄像头画面映射到机器人显示空间的？",
         "核心是先建立两个空间之间的对应关系。视觉模块给出人物或目标在摄像头画面中的位置，机器人端需要把这个位置转换成眼睛在显示区域中应该朝向的位置。我们通过坐标系标定建立这种映射，并在运行中持续更新目标位置。这样人移动时，机器人眼神不会固定在某个点，而是会跟随位置变化。具体映射形式、标定点数量和坐标变换参数，要按我实际实现补充。",
         "不要泛泛地说“做了标定”。准备好解释输入是什么、输出是什么、如何得到映射、映射误差会造成什么现象。"),
        ("检测结果更新后，机器人为什么不会把瞳孔直接跳到新位置？具体是怎样平滑的？",
         "我们把检测坐标和实际渲染坐标分开处理。检测到人的新位置后，系统只更新目标坐标 target_x 和 target_y；屏幕上真正绘制瞳孔的位置 gaze_x 和 gaze_y 不会直接替换为目标值，而是在每一帧逐步靠近。单轴的更新规则是：本帧眼神位置等于上一帧眼神位置，加上“目标位置减去上一帧位置”乘以平滑系数。对应的代码形式是 gaze_x += (target_x - gaze_x) * smoothing，y 轴同理。默认 smoothing 是 0.35，因此每帧移动的是剩余距离的 35%，距离目标越近，单帧移动越小。这个指数平滑过程能过滤检测坐标的高频波动，让眼神呈现自然缓动，而不是机械地左右跳动。",
         "要主动说清两个坐标的职责：target_x/y 来自检测结果，gaze_x/y 才是每帧实际渲染的坐标。"),
        ("为什么在人刚出现和用户远离时，要使用不同的平滑系数？",
         "这里是根据交互状态调整眼神运动速度。当检测到人刚出现并触发到达提示时，系统把平滑系数临时提高到至少 0.70。这样机器人会更快地转向用户，给出及时的“我注意到你了”的反馈；但它仍然遵循逐帧插值，所以不会硬跳转。相反，当用户远离并进入 retreating 状态时，平滑系数会降到最多 0.22，让视线离开的速度更慢。这样既减少了突然转开的生硬感，也让离开的行为更符合陪伴式交互的节奏。",
         "可用一句话概括：到达强调及时回应，因此更快；远离强调自然收束，因此更缓。"),
        ("视觉和语音同时触发时，系统怎样保证反馈不混乱？",
         "我们主要通过并发隔离来保证交互不会互相阻塞。视觉跟踪、眼睛渲染、关键词检测和音频播放分别在独立线程中运行：视觉线程持续更新目标坐标，渲染线程按自己的帧率用平滑后的 gaze_x 和 gaze_y 绘制眼神，关键词检测线程独立等待并识别语音事件，音频线程负责播放反馈声音。这样即使音频播放耗时较长，也不会冻结视觉跟踪和眼神更新。线程之间只传递必要的数据或事件，例如目标坐标和已识别的关键词，避免多个模块直接抢占同一处理流程。",
         "老师若继续追问线程安全，可以回答：共享状态只保留必要字段，并用消息队列或锁保护读写；渲染线程读到最新坐标即可，不需要等待音频播放结束。"),
        ("你如何判断这个系统确实达到了实时交互的效果？",
         "我会从功能链路和用户可感知的响应两层验证。功能上，要分别检查语音触发、本地推理、视觉定位、坐标映射和声音或眼神反馈是否能稳定工作，再检查它们连续运行时是否出现阻塞。体验上，观察用户移动后注视方向是否持续跟随，以及说出关键词后反馈是否足够及时和一致。由于简历没有记录统一的定量指标，我不会虚构准确率或时延；后续可以补充端到端响应时间、关键词识别成功率、跟随误差和连续运行稳定性等指标。",
         "如果你做过演示或日志记录，要在这里加入具体的测试场景与次数。"),
        ("这个项目目前最大的局限是什么？下一步会怎么改进？",
         "当前方案服务的是有限指令和较简单的视觉跟随，因此它的理解能力和交互多样性仍然有限。端侧设备的计算资源也限制了更复杂的感知模型。下一步我会优先补充系统化评估，包括不同光照、距离和背景噪声下的表现；在不明显增加资源占用的前提下，优化目标丢失和坐标抖动时的交互策略；再根据用户反馈扩展更自然的多轮互动。这样改进会先解决可靠性问题，再扩展功能。",
         "局限要真实、可改进。避免回答“没有明显不足”。")
    ]
    for i, (q, a, t) in enumerate(robot_questions, 1):
        add_question(doc, i, q, a, t)

    add_heading(doc, "二 台球击球策略决策项目")
    billiard_questions = [
        ("为什么台球策略适合用蒙特卡洛树搜索，而不是只按几何距离选一杆最容易进的球？",
         "因为台球是连续决策问题。一杆球即使容易打进，也可能让母球停在不利位置，影响下一杆；反过来，有些难度适中的选择会为后续创造更好的局面。只按当前进球难度选球，属于局部判断。蒙特卡洛树搜索会从当前局面出发，尝试不同候选动作，并通过后续模拟估计它们的长期收益，因此更适合比较多步策略。",
         "可以用一个直观例子：眼前最容易进的球可能会让下一颗球被遮挡。"),
        ("为什么选择 MCTS，而不是贪心、穷举搜索、强化学习或其他算法？",
         "选择 MCTS 是因为它和这个项目的约束比较匹配。第一，台球不是单步问题，贪心方法可以快速选出当前最容易进的一杆，但通常无法判断这杆之后母球的位置和连续进攻机会，因此容易陷入局部最优。第二，击球方向、力度等变量是连续的，如果直接做穷举或深度枚举，分支会很快爆炸；我们先用几何约束把连续空间离散成一批可行候选，再让 MCTS 把计算集中在更有希望的分支上。第三，项目已经具备物理仿真环境，MCTS 可以直接通过模拟未来球局来估计长期回报，不需要事先准备大量带标签的最优击球数据。第四，相比从头训练强化学习策略，MCTS 更适合当前以仿真验证和策略比较为目标的阶段：它不需要很长的训练过程，能够在给定计算预算内逐步改进搜索结果，也便于检查每个候选动作、模拟路径和奖励来源。若未来有大量高质量仿真或真实对局数据，强化学习可以作为扩展方向，用来学习候选生成或价值估计；在本项目的初始阶段，MCTS 是在可解释性、数据需求和多步规划能力之间更合适的选择。",
         "回答重点是“项目约束匹配”，不是说 MCTS 一定优于所有算法。可按四点说：贪心看得太短、穷举分支太多、强化学习需要训练数据和训练成本、MCTS 能利用现有仿真做可解释的在线多步搜索。"),
        ("在这个问题里，状态、动作和评价结果分别是什么？",
         "状态至少包括桌面上各球的位置、当前轮次或规则相关信息，以及可能影响后续击球的球局条件。动作是一次击球方案，例如目标球选择、击球方向和力度等连续变量。评价结果不是只看本杆是否进球，而要综合后续可击球机会、得分或胜负状态，以及失误风险。具体状态变量和规则定义应以项目实际仿真环境为准，但这三者必须保持一致，否则搜索结果没有明确含义。",
         "这题非常常见。面试前要能画出一条“状态 -> 动作 -> 仿真 -> 新状态 -> 评价”的流程。"),
        ("连续动作空间会让搜索量很大，你们怎样生成和筛选候选击球方案？",
         "我们没有把所有方向和力度都无差别枚举，而是先结合几何约束生成更有可能成立的方案，例如考虑目标球、袋口、母球和障碍球之间的空间关系，再排除明显无法形成有效路径的动作。这样做相当于先用物理和几何知识缩小候选集合，再把搜索资源留给更值得评估的分支。它不能保证一开始就保留所有理论可能性，但能明显减少无效计算，让后续树搜索更可行。",
         "按真实情况补充你们用了哪些约束，例如路径遮挡、可达角度、边界碰撞或力度范围。"),
        ("蒙特卡洛树搜索在你们的系统中大致怎样运行？",
         "我们的流程从当前球局作为根节点开始。首先，根据球桌几何关系生成一批可行击球方案，例如目标球与袋口连线是否可达、母球到目标球的路径是否被遮挡，以及方向和力度是否在允许范围内；这些方案构成根节点可扩展的动作。随后进行多轮搜索，每一轮包含四步。第一步是选择：从根节点向下，在已经访问过的子节点中选择一个兼顾当前平均收益和探索价值的分支。可以用 UCT 思路衡量，即平均收益加上与父节点访问次数有关的探索项；访问次数少但潜力未知的分支会得到额外探索机会。第二步是扩展：当选中的节点还有未尝试的合法击球方案时，取出其中一个，用物理仿真执行这一杆，得到新的球局，并把它加入树中成为子节点。第三步是模拟：从这个新球局继续用预先定义的默认策略模拟后续若干杆或直到终局，默认策略同样只从几何上可行的候选中选择，以估计这条路线的长期结果。第四步是回传：把本次模拟的评分沿着刚刚经过的路径逐层传回，更新每个节点的访问次数和累计收益。经过固定次数或固定时间预算的迭代后，在根节点选择访问次数最多或平均收益最高的击球方案作为最终决策。这样既不会只盯着一杆是否进球，也不会把全部计算花在一开始看似最优的单一方案上。",
         "面试时可按“候选生成 -> 选择 -> 扩展 -> 模拟 -> 回传 -> 根节点决策”讲。若老师追问 UCT，可说明它由平均收益和探索项组成，用来平衡利用已知好方案与尝试较少访问的方案。"),
        ("物理仿真环境需要考虑哪些因素？怎样保证它有参考价值？",
         "至少要模拟球的运动、球与球之间的碰撞、球与桌边的碰撞，以及进袋和规则判定。仿真的参考价值取决于它是否能合理反映动作与结果的关系，而不只是画面看起来像台球。验证时可以挑选典型球局和简单击球情形，检查运动轨迹、碰撞结果和进袋判定是否符合预期；也可以在统一条件下比较不同策略，而不是只看单次结果。简历中没有列出物理引擎或参数，我会在面试中只说明实际使用过的部分。",
         "老师若追问摩擦、旋转、碰撞恢复系数，应据实回答；没有实现时说明仿真的简化假设及其影响。"),
        ("你们如何设计策略评价？为什么不能只把是否进球作为回报？",
         "我会把策略评价设计成由“终局结果、当前得分、下一杆局面和风险成本”共同构成的回报。首先，赢得球局或完成目标应当有最高优先级；如果没有到终局，则对合法进球和有效得分给正向奖励。其次，要评价母球在这一杆后的落点和下一杆的可行性：例如下一颗目标球是否有清晰路径、击球角度是否合理、母球是否落在容易控制的位置。与此同时，对犯规、母球落袋、无效碰撞、把母球留在明显不利位置等情况施加惩罚。最终可以把这些项按权重组合成一次模拟的总回报，权重由实验调节，但终局结果的权重应高于局部走位。不能只把是否进球作为回报，是因为“进球”只反映这一杆的局部成功。比如一杆球进了，但母球被挡住、下一杆没有可行路径，长期表现可能比一杆难度略高却能留下好走位的方案更差。",
         "可以把评价式口头概括为：终局结果 + 本杆得分 + 下一杆可行性 - 犯规和风险。不要报出具体权重，除非你实际做过调参和对比实验。"),
        ("MCTS 的奖励函数具体是怎样设计的？",
         "奖励函数的目的，是让搜索结果更符合“赢下整局并尽量形成连续进攻”的目标，而不是只追求当前一杆进球。一次模拟的总奖励可以写成：R = R_终局 + R_本杆 + R_走位 + R_下一杆机会 - R_风险。R_终局权重最高：如果完成目标或赢得球局，给予较大正向奖励；如果失败或形成明显不可逆的不利局面，则给予较大惩罚。R_本杆反映当前击球是否有效进球或得分。R_走位评价母球的落点是否便于控制，R_下一杆机会评价下一颗球是否存在清晰、难度合理的进球路径。风险项则惩罚犯规、母球落袋、无效碰撞、目标路径被遮挡，以及把母球留在边库或其他明显不利位置的情况。实际实现中，各项先经过尺度归一化，再通过仿真中不同权重组合的胜率、连续得分能力和计算稳定性来调节权重。",
         "若没有做过系统化权重搜索，应诚实说权重是按任务优先级设定的初始方案，后续需要通过更多球局和消融实验验证。"),
        ("你们如何评估一次击球后的走位和下一杆机会？",
         "走位评价的对象是当前一杆结束后的新球局。首先，系统对每颗仍在桌面上的目标球检查下一杆是否可行：母球到目标球的路径是否被遮挡，目标球到袋口的路径是否清晰，以及母球、目标球和袋口能否形成合理击球角度。满足基本条件的球记为有效候选，候选数量 N_valid 可以反映局面是否有选择余地。其次，我们为每个有效候选计算机会质量 Q_i，主要考虑两段路径的遮挡情况、母球到目标球和目标球到袋口的距离、击球角度与理想进球线的偏差，以及该球打进后是否还能为再下一杆留下路线。走位总分可以概括为：R_position = alpha × N_valid + beta × max(Q_i) + gamma × average(Q_i) - delta × P_bad。这里，max(Q_i) 代表最佳下一杆机会，average(Q_i) 反映整体局面是否宽裕，P_bad 则惩罚母球贴库、被遮挡或没有直接可行进攻路线。这样，系统既偏好保留一个高质量的下一杆，也偏好保留多个可选方案。",
         "面试时先讲“有无机会”，再讲“机会质量”。角度可用母球、目标球、袋口三个位置构成的夹角衡量；角度偏离理想进球线越大，候选评分越低。"),
        ("你们怎样比较不同策略的效果？",
         "我们在统一的仿真环境中重复运行不同策略，使它们面对相同或可比的初始球局，再比较完成任务的表现和策略选择过程。这样可以降低单次随机结果带来的影响。项目已经形成了可重复的对战和性能比较平台；如果要做更完整的评价，我会增加多种难度球局、重复次数、平均得分或成功率、计算时间等指标，并分析策略在哪些情形下更有优势。",
         "如果你有实际实验表格，记住对比对象、测试球局数量、指标和结论。"),
        ("这项工作中最难的部分是什么？你从中学到了什么？",
         "我认为最难的不是把树搜索流程写出来，而是把真实问题转成一个可搜索、可仿真、可比较的决策问题。候选动作太少会遗漏好策略，太多又会让搜索成本过高；回报设计也会直接影响系统偏好。这项工作让我更重视问题建模和实验设计。面对连续决策问题，先明确状态、动作、约束和评价目标，往往比先选一个算法更重要。",
         "用一个你亲自处理过的困难收尾，例如调试碰撞、候选方案异常或评价偏差。")
    ]
    for i, (q, a, t) in enumerate(billiard_questions, 1):
        add_question(doc, i, q, a, t)

    add_heading(doc, "三 人宠互动行为研究")
    pet_questions = [
        ("这个研究为什么要同时使用眼动、视频和生理信息？",
         "人宠互动是一个动态过程，单一数据只能看到其中一部分。视频可以反映人和宠物的位置、距离和行为变化；眼动能反映人当时关注哪里；生理信息可以提供情感或唤醒状态的辅助线索。把它们放在同一时间线上，才有可能分析例如“人在宠物靠近时看向哪里，以及这个阶段是否伴随行为或生理变化”。多模态的价值在于相互补充，而不是简单把数据堆在一起。",
         "不要把生理数据解释得过强。它通常提供关联线索，不能单独证明某种心理原因。"),
        ("视频分析的完整流程是什么？",
         "我会把流程概括为四步：首先从视频中识别人和宠物，获得每一时刻的位置；其次把连续位置连接成二维运动轨迹；然后结合行为编码或自动标注，标出互动中的关键行为片段；最后把这些视频信息与眼动数据按时间对齐，进行跨模态分析。这个流程的结果是把原本连续的视频转成可计算的轨迹、行为标签和时间序列，便于后续统计和比较。",
         "按真实情况补充目标检测、跟踪或标注工具。简历没有模型名时，不要强行报模型。"),
        ("把一段连续的人宠互动转成可计算的时空行为序列，有什么用？",
         "它的作用是把原本只能靠人观看和主观描述的视频，变成可以重复分析和比较的数据。每一时刻都可以记录人和宠物的位置、相对距离、移动速度、是否接近或远离，以及对应的行为标签和人的注视位置。这样，我们可以统计某类行为出现的频率和持续时间，比较不同互动阶段中的距离和运动模式，并进一步分析人在宠物靠近、移动或发生特定行为时是否更常注视它。时空行为序列也使不同参与者、不同视频片段能够在统一指标下比较，并能被后续的统计分析或模型使用。需要强调的是，它提供的是行为和注意力的量化证据，而不是仅凭轨迹直接断言某种情绪或因果关系。",
         "可用一句话概括：从“看视频讲故事”变成“按时间、位置和行为标签进行可重复比较”。"),
        ("人和宠物在画面中发生遮挡，或者宠物短暂离开画面时，你会怎样处理？",
         "我会把“短时遮挡”和“真正离开画面”分开处理。对于短时遮挡，跟踪器不应立刻删除目标，而是保留该目标的 ID、最近位置和速度，用运动预测在后续几帧给出可能位置；新检测结果出现后，再结合位置连续性、框重叠程度和外观特征做匹配，确认它是否仍是同一只宠物。这样可以避免目标被遮挡一下就被分配成新的 ID。对于持续超过设定时长的未检测到，我会结束该段轨迹，并把这段时间标记为“不可观测”或“离开画面”，而不是用预测轨迹长时间补全。后续计算距离、速度、注视目标等指标时，需要排除或单独标注这些缺失片段；对关键片段再做人工复核。这样既尽量保持短时轨迹连续，也不会把算法猜测当成真实行为数据。",
         "这题的关键是说明边界：短时遮挡可以做有限预测和重关联；长时间缺失必须显式标记，不能继续伪造轨迹。"),
        ("眼动数据和视频数据的时间对齐为什么重要，怎么做？",
         "时间对齐是跨模态分析的前提。眼动数据回答的是“人在某个时刻看向哪里”，视频行为数据回答的是“这一时刻人和宠物在做什么、位于哪里”。如果两个时间轴没有对齐，就可能把人注视宠物的记录错误关联到宠物已经离开、接近或做其他动作的片段，之后关于注意力和行为关系的结论也就没有依据。处理时需要先建立共同时间基准，再根据开始事件、同步标记或共同时间戳将眼动记录匹配到对应视频帧和行为标签。对齐后，才能分析不同互动阶段的注视比例、注视持续时间，以及注视变化与距离、运动或行为事件之间的时间关系。",
         "采样率不同可通过重采样、插值或以较低频率为统一时间轴处理；必须报告同步误差，并避免做超过设备时间精度的解释。"),
        ("你们从轨迹和行为标签中可以提取哪些指标？",
         "轨迹可以用于计算人与宠物之间的距离、相对位置、移动速度、停留区域和接近或远离的变化；行为标签可以统计某类行为的出现次数、持续时间和发生顺序。结合眼动后，还可以分析人对宠物或特定区域的注视比例、注视持续时间，以及这些指标在不同互动阶段的变化。选择指标时要先回到研究问题，避免因为能算就全部计算。",
         "准备 2 至 3 个与你们实际研究问题直接相关的指标，并说明它们为什么有意义。"),
        ("自动标注与人工标注各有什么作用？你会怎样验证结果可靠？",
         "自动标注的优势是能处理较长的视频并提高效率，但它可能在遮挡、姿态变化或复杂背景下出错。人工标注能提供更细致的语义判断，也适合作为抽样核查或参考标准。因此，比较稳妥的方式是让自动流程先完成可重复的基础提取，再对关键片段进行人工检查；如果条件允许，可以对比自动标签与人工标签的一致性，并记录常见错误类型。这样既提高效率，也不忽视结果可靠性。",
         "不要声称做过一致性检验，除非你确实做过。可以说这是下一步的严谨验证方案。"),
        ("这项研究的局限和后续方向是什么？",
         "互动行为受个体差异、环境、拍摄角度和实验任务影响很大，因此小样本或单一场景的结论不宜直接推广。视频中的遮挡和数据同步误差也会影响分析。后续可以扩大参与者和互动场景，明确行为标签体系，补充自动结果与人工标注的一致性验证，并在研究设计上把相关关系和因果结论区分清楚。这样的改进会让结果更可靠，也更容易解释。",
         "老师通常看重你是否知道研究结论的边界。")
    ]
    for i, (q, a, t) in enumerate(pet_questions, 1):
        add_question(doc, i, q, a, t)

    add_heading(doc, "四 综合追问与个人贡献")
    combined_questions = [
        ("这三个项目看起来方向不同，它们之间有什么共同主线？",
         "它们的具体应用不同，但都涉及如何从现实世界中获取信息、建立可计算的表示，再据此作出反馈或分析。玩偶机器人强调感知后的实时交互，台球项目强调状态下的连续决策，人宠研究强调多源行为数据的理解。它们让我逐渐形成一个兴趣：把算法和系统放回真实任务中，关注输入是否可靠、决策是否合理，以及结果是否便于验证。",
         "这里不需要硬凑一个很大的研究方向，清楚说明共同的方法意识即可。"),
        ("请说一个项目中遇到的失败或不符合预期的情况，你如何处理？",
         "我会选择自己真实经历过的一件事来回答。回答时先说明现象，例如模型能运行但响应不稳定、候选策略结果不合理或数据对齐出现偏差；再说明如何定位问题，是拆分模块检查、复现输入输出，还是用更简单案例做验证；最后说修复或调整后的结果，以及留下的教训。老师更关注你如何定位和复盘，而不是项目是否从未出错。",
         "请在面试前替换为一个真实案例，并准备“现象、排查、原因、改动、结果”五个要点。"),
        ("如果让你从头重新做其中一个项目，你会怎样改进研究设计？",
         "我会先把问题定义和评价标准定得更早、更明确。例如机器人项目可以先定义响应时间、跟随稳定性和交互成功率；台球项目可以在开始前确定典型球局和策略比较指标；人宠项目可以先确定标签体系、同步方案和数据质量检查。这样能让后续方法选择直接服务于要验证的问题，也能减少做到后期才发现无法比较的情况。",
         "这个回答体现研究设计能力，注意结合一个你最熟悉的项目举例。"),
        ("你在团队项目中如何说明自己的贡献，避免把团队成果全部说成个人成果？",
         "我会先用“我们”讲项目目标和总体成果，再用“我主要负责”讲自己的具体模块、决策或排查工作。例如，在玩偶机器人项目中，我可以明确说明自己参与了哪些端侧部署、映射或流程衔接工作；对于其他成员负责的部分，也会如实说明。这样既能解释自己掌握的技术细节，也能体现我理解团队协作边界。",
         "面试中最忌讳含糊地说“我都做了”。准备每个项目 2 至 3 项可核实的个人工作。"),
        ("你的论文在投，面试老师问你论文的核心问题和当前状态，你会怎样回答？",
         "我的论文关注参考生成的机器人眼部动画中，角色身份连续性是否能够保持。核心问题是，生成出的眼部动画虽然可能有表现力，但不同片段或不同参考条件下，角色是否仍然让人感到是同一个角色，需要系统评估。当前论文已投稿 CHI 2027，处于在投状态。介绍时我会清楚区分已完成的工作和仍在评审中的结论，不把在投论文说成已发表成果。",
         "若老师追问实验设计、数据或贡献，请只回答你确实参与和熟悉的部分。")
    ]
    for i, (q, a, t) in enumerate(combined_questions, 1):
        add_question(doc, i, q, a, t)

    doc.add_page_break()
    add_heading(doc, "五 英语项目问答")
    p = doc.add_paragraph()
    set_para_format(p, after=9, line=1.35)
    r = p.add_run("练习原则：")
    set_run_font(r, size=10.5, bold=True, color=(35, 78, 121))
    r = p.add_run("英语回答先用一句话说清“我用了什么方法，解决了什么问题”，再补充原因和结果。语速放慢，尽量使用短句。")
    set_run_font(r, size=10.5)

    english_questions = [
        ("Could you introduce the method you used in the doll robot interaction project?",
         "In this project, we combined visual gaze tracking with fast keyword detection. The goal was to let a small doll robot notice a person's position and respond to simple voice commands. Because the target device had limited computing power, we used lightweight models and deployed them locally on a Raspberry Pi Zero 2 W. We mapped camera coordinates to robot display coordinates. To avoid abrupt gaze movement, we kept the detected target position separate from the rendered gaze position and updated the gaze frame by frame with exponential smoothing. The system could provide local, real-time interaction without relying on a network connection.",
         "在这个项目中，我们把视觉注视跟随和快速关键词检测结合起来。目标是让一个小型玩偶机器人能够感知人的位置，并对简单语音指令作出回应。由于目标设备的计算能力有限，我们使用了轻量级模型，并将它们本地部署在树莓派 Zero 2 W 上。我们把摄像头坐标映射到机器人显示坐标。为避免眼神突然跳动，我们将检测到的目标位置与实际渲染的眼神位置分开，并通过逐帧指数平滑更新眼神位置。该系统不依赖网络，也能实现本地实时交互。",
         "lightweight models；local deployment；coordinate mapping；exponential smoothing；real-time interaction"),
        ("Why did you use keyword detection and INT8 quantization on the robot?",
         "We used keyword detection because the robot only needed to recognize a small set of voice triggers. This was more suitable for the device than a large and open-ended speech recognition system. INT8 quantization reduced the memory and computing cost of the model, which made local inference more practical on the Raspberry Pi. The main trade-off was that we had to balance model efficiency with recognition performance. Our purpose was to make the interaction stable and responsive on limited hardware.",
         "我们使用关键词检测，是因为机器人只需要识别少量的语音触发词。这比大型、开放式的语音识别系统更适合该设备。INT8 量化降低了模型的内存和计算开销，使得模型更容易在树莓派上进行本地推理。主要的权衡是模型效率与识别效果之间的平衡。我们的目的是让有限硬件上的交互保持稳定且响应及时。",
         "voice triggers；computing cost；local inference；trade-off"),
        ("What problem did Monte Carlo Tree Search solve in your billiards project?",
         "The billiards task was a long-term decision-making problem. A shot that looked good at the current moment might create a poor position for the next shot. Monte Carlo Tree Search helped us compare candidate shots by simulating possible future outcomes. Instead of choosing a shot only by its immediate success, the method considered future opportunities and risks. This made the decision process more suitable for multi-step planning.",
         "台球任务是一个长期决策问题。当前看起来很好的击球，可能会给下一杆造成很差的位置。蒙特卡洛树搜索帮助我们通过模拟未来可能的结果来比较候选击球方案。该方法不只根据一杆是否成功来选球，还考虑后续机会和风险。因此，它更适合多步规划。",
         "long-term decision-making；candidate shots；future outcomes；multi-step planning"),
        ("How did you reduce the search space in the billiards project?",
         "The action space in billiards is continuous, so searching every possible direction and force would be too expensive. We used geometric constraints to generate and filter candidate shots before the tree search. For example, we considered the spatial relationship among the cue ball, target balls, pockets, and possible obstacles. This removed many clearly infeasible shots and allowed the search to focus on more promising options. The trade-off was that the constraints needed to be carefully designed so that useful candidates were not removed too early.",
         "台球的动作空间是连续的，因此搜索每一种可能的方向和力度成本太高。我们在树搜索之前使用几何约束来生成和筛选候选击球方案。例如，我们考虑母球、目标球、袋口和可能障碍物之间的空间关系。这样能排除许多明显不可行的击球，把搜索资源集中在更有希望的方案上。相应的权衡是，约束需要谨慎设计，避免过早排除有价值的候选方案。",
         "continuous action space；geometric constraints；filter candidates；infeasible shots"),
        ("What was the purpose of combining eye tracking and video analysis in the human-pet interaction study?",
         "The purpose was to understand the interaction process from more than one perspective. Video analysis provided the positions and movement trajectories of people and pets. Eye tracking showed where the person was looking during the interaction. By aligning these data over time, we could study how visual attention changed with behavior and spatial movement. This reduced the need to rely only on manual frame-by-frame observation and made the interaction process more measurable.",
         "这样做的目的是从多个角度理解互动过程。视频分析提供人和宠物的位置以及运动轨迹，眼动数据反映人在互动时看向哪里。通过在时间上对齐这些数据，我们可以研究视觉注意力如何随行为和空间移动而变化。这样能够减少只依靠人工逐帧观察的局限，并使互动过程更便于量化分析。",
         "movement trajectories；visual attention；align data over time；measurable"),
        ("How did you align different types of data in the human-pet interaction study?",
         "We treated time alignment as a necessary step before cross-modal analysis. The video, eye-tracking, and physiological data had to be connected to a common time reference. Then we could match a gaze record with the corresponding video frame and behavior label. This was important because, without time alignment, we could not reliably say what the participant was looking at during a specific interaction event. The exact synchronization method depends on the recording setup, so I would explain the method we actually used in the experiment.",
         "我们把时间对齐视为跨模态分析之前必须完成的一步。视频、眼动和生理数据需要连接到共同的时间参考上。之后，我们才能把一条注视记录对应到相应的视频帧和行为标签。这很重要，因为如果没有时间对齐，我们就无法可靠地说明参与者在某个具体互动事件中看向哪里。具体的同步方法取决于采集设置，因此我会根据实验中实际使用的方法进行说明。",
         "common time reference；cross-modal analysis；corresponding frame；synchronization"),
        ("What did you learn from these projects?",
         "I learned that an algorithm should be chosen based on the real problem and the constraints of the system. In the robot project, the main constraint was limited hardware, so efficient local inference was important. In the billiards project, the key issue was long-term planning, so we needed to evaluate future outcomes. In the human-pet study, the challenge was understanding a complex interaction process, so we combined multiple types of data. These projects taught me to connect methods, evaluation, and practical requirements.",
         "我认识到，算法选择应该基于真实问题和系统约束。在机器人项目中，主要限制是硬件能力有限，因此高效的本地推理很重要。在台球项目中，核心问题是长期规划，因此需要评估未来结果。在人宠研究中，难点是理解复杂互动过程，因此我们结合了多类数据。这些项目让我学会把方法选择、效果验证和实际需求联系起来。",
         "system constraints；long-term planning；multiple data types；evaluation")
    ]
    for i, item in enumerate(english_questions, 1):
        add_english_question(doc, i, *item)

    add_heading(doc, "六 面试前最后核对")
    final_checks = [
        "为每个项目准备一句话版本、60 秒版本和 2 分钟版本。",
        "把文中“按真实情况补充”的位置替换为自己实际使用的工具、参数、实验场景或排查经历。",
        "核对团队贡献边界。项目总体成果用“我们”，个人工作用“我主要负责”或“我参与了”。",
        "对于没有做过的技术细节，说明当前方案的边界和你会如何验证，不要猜测或编造。",
        "英语回答重点练习加粗概念对应的发音，并确保能够脱稿说清方法、问题和效果。"
    ]
    for check in final_checks:
        p = doc.add_paragraph(style="List Bullet")
        set_para_format(p, after=4, line=1.35)
        r = p.add_run(check)
        set_run_font(r, size=10.5)

    doc.add_page_break()
    add_heading(doc, "七 科研项目英文介绍")
    p = doc.add_paragraph()
    set_para_format(p, after=8, line=1.35)
    r = p.add_run("使用建议：")
    set_run_font(r, size=10.5, bold=True, color=(35, 78, 121))
    r = p.add_run("以下每段约 45 至 60 秒，适合老师要求用英语介绍项目时使用。先记住每段的目标、问题、方法和结果，再用自己的语言表达。")
    set_run_font(r, size=10.5)

    project_intros = [
        ("项目一 玩偶机器人多模态交互系统",
         "My project focused on a doll robot that can respond to people through gaze and voice. The main challenge was that the robot used a Raspberry Pi Zero 2 W, which has limited computing resources. We therefore deployed lightweight keyword detection and visual tracking locally, so the robot did not need a network connection for basic interaction. The vision module updates a target position, while the rendering module smoothly moves the displayed gaze toward that target frame by frame. Visual tracking, rendering, keyword detection, and audio playback run in separate threads, so audio playback does not interrupt gaze following. As a result, the robot can follow a person's position and provide timely eye and voice feedback on the device.",
         "我的项目聚焦于一个能够通过眼神和语音与人互动的玩偶机器人。主要挑战是机器人使用树莓派 Zero 2 W，计算资源有限。因此，我们将轻量级关键词检测和视觉跟踪部署在本地，让机器人在基本交互中不依赖网络。视觉模块更新目标位置，渲染模块则逐帧让显示的眼神平滑靠近该目标。视觉跟踪、眼睛渲染、关键词检测和音频播放在独立线程中运行，因此音频播放不会中断眼神跟随。最终，机器人能够在设备端跟随人的位置，并及时给出眼神和语音反馈。"),
        ("项目二 基于蒙特卡洛树搜索的台球击球策略决策",
         "This project studied shot selection in billiards as a multi-step decision-making problem. A shot that is easy to make now may leave a poor position for the next shot, so we could not evaluate actions only by immediate success. We first generated feasible candidate shots using geometric constraints, such as line-of-sight and pocket paths. Then we used Monte Carlo Tree Search to select, expand, simulate, and back-propagate candidate actions in a physics simulation environment. The reward considered the final game outcome, current scoring, cue-ball position, future shot opportunities, and risks such as fouls. This allowed the system to compare long-term value instead of choosing only the most obvious current shot.",
         "这个项目把台球击球选择视为一个多步决策问题。当前容易打进的一杆，可能会给下一杆留下很差的位置，因此不能只用眼前是否进球来评价动作。我们首先利用视线、进袋路径等几何约束生成可行候选击球方案，再在物理仿真环境中使用蒙特卡洛树搜索，对候选动作进行选择、扩展、模拟和回传。奖励函数同时考虑终局结果、当前得分、母球走位、未来击球机会，以及犯规等风险。这使系统能够比较长期价值，而不是只选择当前看起来最直接的一杆。"),
        ("项目三 基于眼动和视频分析的人宠互动行为研究",
         "In this project, we built a multimodal temporal analysis framework that combines eye-tracking data with video data. I was mainly responsible for the data-processing side. We used object detection to extract the spatiotemporal positions of people and pets and constructed two-dimensional movement trajectories. Then, by combining behavior coding with automatic annotation, we aligned video-based behaviors with eye-tracking data over time. This process translates a continuous interaction into a computable spatiotemporal behavior sequence, which supports repeatable analysis of movement, distance, behavior events, and visual attention across interaction stages.",
         "在这个项目中，我们构建了融合眼动和视频的多模态时序分析框架。我主要负责数据处理部分：用目标检测提取人和宠物的时空位置、构建二维运动轨迹，再结合行为编码与自动标注，把视频行为和眼动数据在时间上对齐。这个过程将一段连续互动转化为可计算的时空行为序列，从而支持对不同互动阶段中的移动、距离、行为事件和视觉注意力进行可重复分析。")
    ]
    for heading, english, chinese in project_intros:
        add_heading(doc, heading, level=2)
        add_label_paragraph(doc, "English introduction: ", english)
        add_label_paragraph(doc, "中文翻译：", chinese)

    add_heading(doc, "八 在投论文面试准备")
    add_label_paragraph(doc, "论文题目：", "Expressive Eyes Consistent Characters Evaluating Identity Continuity in Reference Generated Robot Eye Animations。论文当前为 CHI 2027 在投稿件，应表述为“在投”或“已投稿”，不要说成已发表。")
    add_label_paragraph(doc, "论文一句话：", "论文研究如何把人物、卡通或动物的眼部参考图转化为机器人可执行的眼睛样式，并检验眼睛在切换表情后是否仍然让观察者觉得是同一个角色。")
    add_question(doc, "论文", "老师您好，这篇论文主要是关于什么的？",
         "这篇论文关注参考图生成的机器人眼睛在动画中的身份连续性。很多方法可以根据一张人物、卡通或动物的参考图生成一对看起来相似的中性眼睛，但机器人眼睛还需要眨眼、改变眼睑形状和表达情绪。此时，原本能识别角色的瞳孔比例、虹膜颜色、眼睛轮廓和高光等特征可能会被改变或遮挡，最后看起来像另一种角色。我们提出 EyeStyleMimic，把参考图中的关键视觉线索整理成一个由 21 个字段组成、能被固定渲染器执行的眼睛样式。具体流程是先由视觉语言模型提取与渲染有关的特征，再由语言模型把这些观察结果转化为满足参数范围和格式约束的样式值，最后用同一套样式渲染中性眼睛和不同表情。我们用 36 张参考图，包括人物、卡通和动物，并邀请 26 名参与者完成配对比较。参与者先判断哪一对中性眼睛更像参考角色，再比较开心、悲伤、愤怒和疲惫四种动画状态下，哪一组仍更像同一个角色。在 936 组配对比较中，EyeStyleMimic 相对三种实现的基线方法，在静态参考匹配和动态身份连续性上都获得了更高的被选择比例，大约在 76% 到 81.7% 之间。论文的结论是，在本研究的固定渲染器和对照条件下，这条完整流程更容易保留角色可识别性；它并不证明适用于所有眼睛风格、所有模型或所有机器人平台。",
         "建议把这段控制在 90 秒左右。先讲研究问题，再讲方法、实验、结果和边界，老师容易判断你是否真正理解论文。"),
    add_question(doc, "论文", "论文中的 EyeStyleMimic 和直接用计算机视觉或视觉语言模型有什么区别？",
         "区别不只是使用了模型，而是是否把参考信息组织成可稳定执行的样式表示。直接图像测量可以得到颜色和粗略几何信息，但可能把背景、皮肤或阴影误当成眼部特征；视觉语言模型可以描述“瞳孔较大”或“有暖色渐变”等语义特征，但这些描述不一定能直接转换成合法的渲染参数。EyeStyleMimic 先保留视觉模型对参考图的结构化观察，再让语言模型根据渲染器的字段定义、取值范围和坐标约定，生成一套完整的 21 项参数，并与固定基础样式合并。这样输出不是一张静态图片，而是一套能在不同表情下复用的样式。论文比较了基于图像测量、基于视觉语言描述、直接组合两类信息，以及 EyeStyleMimic 四条具体实现路径。",
         "不要说 EyeStyleMimic 已经解决了所有多源信息冲突。论文明确指出，不同信息源的可靠性校准和冲突处理仍是后续问题。"),
    add_question(doc, "论文", "你们如何评估“身份连续性”，而不是只评估静态图片像不像？",
         "我们把表情变化当成对角色身份的压力测试。每个候选眼睛样式先在中性状态下与参考图比较，随后用同一套样式生成开心、悲伤、愤怒和疲惫四种动画状态。参与者不是给单个表情打分，也不是判断情绪是否识别正确，而是在两组动画之间选择哪一组在表情变化后仍更像参考图中的同一个角色。这个任务将静态参考匹配和动态身份连续性区分开来：前者问中性眼睛是否像，后者问经过眼睑、轮廓和可见虹膜区域变化后，角色特征是否仍被保留。",
         "关键术语是 identity continuity，不是 emotion recognition。论文不以“表情识别准确率”作为主要结论。"),
    add_question(doc, "论文", "作为二作，你怎样介绍自己在论文中的角色？",
         "我会先明确这是团队合作成果，并把总体方法、实验和结论用“我们”来表述。接着，我会只说明自己实际参与的工作，例如参考材料整理、机器人眼睛系统实现、实验材料生成、数据处理、用户研究执行、结果可视化或论文撰写中的具体部分。一个稳妥的表达是：‘我作为二作，主要参与了【请按实际分工填写具体模块】；这部分工作让我熟悉了从参考图到可执行眼睛样式的流程，以及如何通过用户配对比较评估动态身份连续性。’如果老师继续追问，我会展开我亲自完成的模块，并如实说明其他成员负责的部分。",
         "请务必把方括号中的内容替换为真实贡献。论文作者目前匿名，且你未在本次材料中说明具体分工，因此文档不替你虚构个人贡献。"),
    add_question(doc, "论文", "这篇论文有哪些局限？",
         "第一，研究只在一套固定渲染器、36 张参考图和三条具体对照路径下进行，因此不能把结果推广为所有计算机视觉、所有视觉语言模型或所有机器人眼睛系统的普遍结论。第二，参与者为 26 名 20 至 28 岁的大学生，样本规模和人群范围有限。第三，实验采用两两比较，测量的是相对选择而非单个系统在真实部署中的绝对满意度。第四，论文中的证据驱动校正器只在 12 张校准图上审查过，目前没有产生视觉改进案例，因此不能声称它已经提升生成质量。后续可以增加更丰富的参考类型和用户群体，加入专家标注或真实部署研究，并验证哪些参数修正真的能改善身份连续性。",
         "讲局限并不会削弱论文，反而体现你对证据范围和后续研究方向有清楚认识。")

    doc.add_page_break()
    add_heading(doc, "九 博士资格面试补充问答")
    add_label_paragraph(doc, "面试侧重点：", "博士资格面试不仅考察项目是否完成，更关注研究问题是否清楚、方法选择是否有依据、证据是否支持结论，以及申请人是否具备独立推进研究的潜力。以下问题与前文的项目技术细节问答配套使用。")
    add_heading(doc, "自我介绍表述校正", level=2)
    add_question(doc, "校正", "你在介绍中把树莓派 Zero 2 W 称为微控制器级别，这个说法准确吗？",
         "不够准确。树莓派 Zero 2 W 是可以运行 Linux 的低算力单板计算机，不是微控制器。更严谨的说法是：它的 CPU、内存和功耗都受到明显限制，无法直接运行桌面端或云端常用的较重模型，因此项目需要在本地推理、模型大小和响应速度之间做取舍。",
         "自我介绍中建议直接使用“低算力单板计算机”或“资源受限的端侧平台”。"),
    add_question(doc, "校正", "人宠互动时的情绪能量化吗？这句话是否过强？",
         "这句话需要更谨慎。视频、眼动和生理数据可以量化人宠互动中的注意分配、空间行为、行为节奏，以及与情绪或唤醒相关的线索，但通常不能仅凭这些数据直接确认某一种情绪。更严谨的研究表述是：我们分析互动过程中的行为、注意力和生理变化，并研究它们之间的时间关联。若要对情绪做更强的判断，还需要明确的量表、自报告或经过验证的情绪标注。",
         "自我介绍可改为“人宠互动中的注意、行为和情绪相关线索能否被量化分析”。"),

    add_heading(doc, "研究主线与博士规划", level=2)
    doctoral_questions = [
        ("三个项目分别解决了感知、决策和行为分析问题。它们之间有什么真正的研究主线？",
         "我把它们放在“从感知到理解再到反馈”的链条中理解。玩偶机器人关注如何在资源受限设备上实时感知用户并给出反馈；台球项目关注系统在状态变化中如何评价长期后果并作出决策；人宠研究关注如何从多模态数据中把真实互动转化为可分析的行为表示。它们共同训练了我从输入可靠性、状态表示、决策机制和结果验证四个方面看待智能系统。博士阶段，我希望将重点收敛到低算力具身智能体如何把多模态感知转化为稳定、自然且可解释的交互反馈。",
         "回答时不要把三个项目说成同一种技术，而要说明它们在同一个系统链条中承担不同角色。"),
        ("你所说的“自然反馈”具体是什么？如何评价？",
         "自然反馈至少包含连续性、及时性和情境匹配三方面。连续性是眼神、动作或表情不会因检测噪声而突跳；及时性是用户发出动作或语音后，系统能在可接受的时间内作出反应；情境匹配是系统反馈与用户位置、语音事件和当前交互状态一致。评价上，我会结合客观指标和主观感受：例如端到端响应时间、眼神轨迹抖动、目标丢失后的恢复时间、关键词误触发率，以及用户对“是否被注意到”“反馈是否自然”的主观评分或配对选择。",
         "避免只说“用户觉得自然”。需要说明可观测指标和用户研究如何共同支撑结论。"),
        ("你的项目看起来偏工程实现。怎样把它发展成可发表的研究问题？",
         "工程原型首先回答“系统能不能做出来”，研究还要回答“为什么这样设计”“相对什么方案更好”“在什么条件下有效”。以玩偶机器人为例，我可以把问题从部署模块转化为：不同感知噪声、响应延迟和反馈平滑策略，怎样影响用户对机器人关注感和自然度的判断。之后设置对照条件，例如无平滑、固定平滑、状态自适应平滑，定义客观指标和用户评价，再用重复实验检验假设。这样系统实现就成为研究变量和验证平台，而不是研究结论本身。",
         "这是博士资格面试中的核心回答。建议用一个熟悉项目讲清“假设、条件、指标、结论边界”。"),
        ("博士阶段如果只能选一个具体问题持续研究，你会选什么？",
         "我希望研究资源受限的具身智能体如何利用多模态感知形成稳定、可解释、对人友好的交互状态。具体来说，我关心视觉、语音和行为线索存在噪声或缺失时，系统怎样判断应该关注什么、何时回应、用什么强度回应，以及这种策略是否真的让用户感到被理解。这个问题既需要系统实现，也可以通过感知鲁棒性、响应时序和用户研究进行验证。",
         "方向要足够聚焦，避免回答“人工智能、人机交互和机器人我都想做”。"),
        ("你目前最大的研究短板是什么？准备怎样补足？",
         "我目前的优势是把算法和系统放到具体任务中实现和调试，但还需要进一步加强严格实验设计、统计推断和研究问题收敛能力。博士阶段，我会通过复现领域代表工作、在项目早期明确假设和评价指标、设计消融实验、学习用户研究与统计分析方法来补足。同时，我会坚持记录失败案例和分析过程，避免只展示系统成功运行的结果。",
         "不要回答“没有明显短板”。能够指出短板并提出具体补足路径，通常更有说服力。"),
        ("你如何区分一个能运行的工程原型和一项值得发表的研究工作？",
         "工程原型证明一个功能可以运行；研究工作还必须提出清楚的问题，说明方法选择的依据，并通过与基线的比较回答它在哪些条件下有效、效果为何出现以及结论的边界。例如论文中的机器人眼睛系统本身只是原型，但通过把动态表情变化定义为身份连续性的测试，并用用户配对研究比较不同生成路径，才形成了可检验的研究贡献。",
         "回答应包含问题、比较、证据和边界四个关键词。"),
        ("如果实验结果与原先假设相反，你会怎样处理？",
         "我会先检查实现、数据质量、随机化和分析流程，确认结果不是由 bug、样本异常或指标错误造成的。若结果可靠，我不会为了得到预期结论而不断修改指标或筛选数据，而是分析假设在哪些条件下不成立，检查是否存在未考虑的机制或交互因素。之后如实报告负结果和适用边界，并把它作为下一轮研究设计的依据。",
         "这题主要考察科研诚信和复盘能力。")
    ]
    for i, (q, a, t) in enumerate(doctoral_questions, 1):
        add_question(doc, f"主线{i}", q, a, t)

    add_heading(doc, "玩偶机器人项目追问", level=2)
    robot_doctoral_questions = [
        ("你说负责端侧部署与推理优化，哪些工作是你亲自完成的？",
         "我会先用“我们”介绍项目的整体目标，再明确我的模块边界。我主要负责的是语音关键词检测和视觉相关模块在端侧的部署、推理流程调试，以及视觉目标坐标到显示端眼神反馈之间的衔接和优化。对于模型训练、硬件结构、机械设计或其他成员负责的部分，我会如实区分。这样既能说明我掌握的技术细节，也不会把团队成果都归为个人成果。",
         "请把“我主要负责”的具体模块替换成真实分工，并为每个模块准备一个调试或优化实例。"),
        ("为什么选择关键词检测，而不使用完整语音识别或大模型语音接口？",
         "当前交互目标是识别有限的语音触发词，因此关键词检测已经能覆盖核心需求。它的模型更小、推理更快、内存和功耗开销更低，也更适合离线运行。完整语音识别或大模型语音接口能够支持更开放的对话，但会带来更高的计算需求、网络依赖和响应不确定性。未来如果要扩展多轮对话，可以将关键词检测保留为低延迟唤醒层，再按设备能力接入更复杂的语音理解模块。",
         "回答重点是任务需求和资源约束，而不是笼统说大模型“不好”。"),
        ("INT8 量化带来什么收益？你如何证明收益不是主观感受？",
         "INT8 量化的直接作用是降低模型参数和计算的表示精度，从而减少存储、内存占用和推理开销。要证明收益，需要在相同硬件、相同输入和相同运行条件下比较量化前后的模型文件大小、峰值内存、平均或分位推理时间，以及关键词识别成功率和误触发情况。只有同时报告效率变化和功能损失，才能说明量化是否值得。若目前没有完整测量，我会将这些指标列为下一步的系统评估计划，而不会只说“它更快”。",
         "准备好按实际数据回答。没有数据时，不要虚构速度提升比例。"),
        ("关键词检测如何应对误触发、背景噪声和多人说话？",
         "我会从触发前、触发时和触发后三层降低误触发。触发前可用语音活动检测、目标在场状态或环境噪声门限限制无关音频进入检测器；触发时可设置置信度阈值和连续帧确认，避免单个瞬时高分直接触发；触发后设置短暂冷却时间或事件去重，避免同一段声音连续触发。还应记录误触发场景和置信度，针对噪声、相似发音和多人说话分别测试。",
         "如果以上策略尚未全部实现，应说“可采用”或“下一步会加入”，不要说成已完成。"),
        ("你如何验证坐标标定和注视跟随真的有效？",
         "我会将验证分为几何精度和交互体验两部分。几何精度上，在摄像头画面中选取多个已知位置，让目标分别出现在中心、边缘和不同距离处，比较映射后的注视点与期望显示位置之间的偏差。交互体验上，观察用户移动时机器人是否持续看向用户、是否出现明显滞后或抖动，并可让用户对自然度和被关注感进行评分。这样可以区分“坐标映射准确”与“用户感觉自然”这两类结果。",
         "若使用多点标定或拟合映射，请按实际方法补充。"),
        ("并发线程之间如何避免数据竞争和过期事件？",
         "我的原则是让每个线程承担单一职责，并尽量减少共享可变状态。视觉线程只更新最新目标坐标，渲染线程周期性读取坐标并平滑绘制；关键词检测线程输出离散语音事件，音频线程消费事件并播放声音。共享坐标和状态通过锁、原子变量或消息队列保护；语音事件附带时间戳或状态，过期事件可以丢弃。渲染线程读取最新状态即可，不需要等待音频播放完成，因此长音频不会卡住眼神反馈。",
         "按真实实现说明使用的是锁、队列还是其他同步机制。"),
        ("目标突然丢失时，为什么不立即把眼神恢复到默认位置？",
         "短暂丢失可能来自遮挡、检测抖动或单帧失败。若立即回到默认位置，机器人会显得注意力不稳定。因此可以在短时间内保持最后一个稳定目标，或者让眼神缓慢回归默认位置；超过设定超时后再进入未检测到目标的状态。超时时间不应只凭直觉设定，应比较不同阈值下的错误注视时长、恢复速度和用户自然度评价。",
         "这题体现状态估计和交互连续性的意识。")
    ]
    for i, (q, a, t) in enumerate(robot_doctoral_questions, 1):
        add_question(doc, f"机器人{i}", q, a, t)

    add_heading(doc, "台球决策项目追问", level=2)
    billiard_doctoral_questions = [
        ("你提到五维连续动作空间，五个维度分别是什么？",
         "这题必须按项目实际的动作定义回答。我会先区分离散决策和连续控制：例如目标球或袋口的选择通常是离散变量，而球杆方向、力度、击打点和旋转相关参数可能构成连续部分。面试时我不会只说“五维”，而会逐项说明每一维的物理意义、范围和是否被离散化。若当前实现实际上没有完整建模旋转或击打点，也会明确说明简化假设及其影响。",
         "请在面试前把五个维度写成准确列表；这一题不能用泛泛答案代替。"),
        ("几何剪枝怎样避免删掉真正有价值的策略？",
         "剪枝一定会在效率和完备性之间取舍，因此我只会删除违反规则、路径被明确遮挡或无法形成有效进袋路径的动作。对于边界较模糊、可能需要反弹或高难度走位的动作，不应简单删除，而可以保留为低优先级候选。之后可通过放宽约束的消融实验，比较候选覆盖率、搜索时间和最终收益，检查剪枝是否频繁遗漏高价值方案。",
         "核心观点：只剪明显不可行的动作，对不确定动作降优先级而不是一刀切。"),
        ("物理仿真中包含哪些因素？哪些因素被简化？",
         "仿真至少需要包含球的运动、球与球碰撞、球与库边碰撞、进袋以及规则判定。更复杂的因素包括摩擦、碰撞恢复系数、旋转、塞和球杆击打位置等。我的回答必须以实际实现为准：已建模的因素说明如何验证，未建模的因素说明它们可能使真实击球与仿真结果产生偏差。因此，当前结论应限定为统一简化物理环境下的策略比较，而不是直接等同于真实台球比赛表现。",
         "不要声称实现了没有实现的旋转、摩擦或真实球杆动力学。"),
        ("奖励函数中的各项怎样归一化，为什么可以加权相加？",
         "终局结果、得分、走位和风险的数值范围不同，不能直接相加。实现时应先把各项映射到统一范围，例如 [0, 1] 或 [-1, 1]，再按任务优先级设置权重，其中终局结果通常权重最高。之后在固定球局集合上进行敏感性分析或消融实验，比较不同权重组合对任务完成率、连续得分、犯规率和决策时间的影响。这样可以避免某个数值较大的项在无意中主导搜索。",
         "若尚未完成权重搜索，可说明这是需要进一步验证的模型设计环节。"),
        ("如何证明性能提升来自 MCTS，而不是候选生成或奖励函数？",
         "需要设计消融和基线对照，而不是只报告最终系统表现。可比较只用几何候选加贪心选择、没有几何剪枝的 MCTS、移除走位与下一杆奖励的 MCTS、不同搜索预算下的 MCTS，以及随机可行击球策略。若完整方法在相同初始球局中同时提高完成率或累计回报，并且消融某一部分后性能下降，才能更有依据地说明各模块的贡献。",
         "这题考察研究设计能力。回答中应出现“相同初始球局、明确基线、消融实验、多个指标”。"),
        ("你会怎样设置策略比较的基线和指标？",
         "基线可以包括随机可行击球、仅按当前进球难度选择的贪心策略，以及不同预算或不同剪枝强度的 MCTS。评价不能只看单次是否进球，还应包括任务完成率、平均累计回报、连续得分次数、犯规率、平均决策时间和不同球局难度下的稳定性。所有策略应在相同的初始球局和随机条件下重复运行，避免把偶然结果误认为方法差异。",
         "准备一套你实际能够运行的基线和指标，避免列出目前无法验证的内容。")
    ]
    for i, (q, a, t) in enumerate(billiard_doctoral_questions, 1):
        add_question(doc, f"台球{i}", q, a, t)

    add_heading(doc, "人宠研究与论文追问", level=2)
    research_paper_questions = [
        ("人宠互动研究具体想回答什么科学问题？",
         "项目不应只停留在“把视频和眼动对齐”。更具体的问题可以是：在人宠互动的不同阶段，人的注视分配是否会随着人宠距离、相对运动和行为类型发生系统变化；这些变化是否与生理线索存在时间关联。这样的表述将数据处理流程与可检验的研究问题连接起来，也明确了当前研究主要分析关联而不是直接证明因果。",
         "请根据真实实验任务和数据范围，把示例问题进一步收敛。"),
        ("你使用哪些生理指标？它们究竟能说明什么？",
         "我会按实际采集到的生理信号说明，例如心率、皮电或其他指标。它们通常反映与唤醒、压力或注意投入相关的变化，不能脱离任务和其他数据直接被解释为某种确定情绪。因此，我会把生理数据视为辅助线索，并与视频行为、眼动位置和实验阶段联合分析。",
         "必须用真实采集的指标替换示例，不能笼统地说“生理数据能判断情绪”。"),
        ("不同设备的采样率不同，时间对齐误差会怎样影响结论？",
         "不同采样率的数据需要通过共同时间戳、同步事件或起始标记对齐，再按分析目标重采样或插值到统一时间轴。如果对齐存在偏差，一条眼动记录可能被关联到错误的视频行为片段，进而改变对注意力和互动关系的判断。因此需要记录同步方式、估计误差，在关键事件附近进行人工核查，并避免对时间分辨率高于设备能力的结论作过强解释。",
         "这题考察你是否理解数据预处理会直接影响研究结论。"),
        ("自动检测和标注的误差如何避免被当成真实行为差异？",
         "我会保留检测置信度，标记遮挡和缺失区间，并对关键片段进行抽样人工复核。对于置信度较低或目标长期缺失的片段，不应直接纳入精细轨迹或注视目标统计，而应排除、单独标记或做敏感性分析。若条件允许，还应比较自动标注与人工标注的一致性，并记录常见错误类型。这样可以把算法误差和真实行为变化尽量区分开。",
         "不要把自动检测输出直接当作无误差的研究真值。"),
        ("如果人与宠物距离变近，同时注视比例变高，能否说宠物靠近导致人更关注它？",
         "不能直接得出因果结论。这两种变化可能共同受到互动任务、声音刺激、个体差异或场景因素影响。当前多模态时序分析更适合报告时间上的同步变化或统计关联。若要研究因果，需要设计可控实验，例如操纵宠物接近行为或互动条件，并控制其他影响因素。",
         "这是研究推断边界题。关键词是“关联不等于因果”。"),
        ("作为二作，论文中你怎样清楚说明自己的贡献？",
         "我会先强调论文是团队合作，并用“我们”表述总体方法、实验和结论。随后只说明自己实际参与的工作，例如机器人眼睛渲染与部署、实验材料生成、用户研究执行、数据处理、图表制作或论文写作中的具体模块。一个稳妥表述是：‘我作为二作，主要负责【真实模块】。这部分工作让我熟悉了从参考图到可执行眼睛样式的流程，以及如何用用户配对比较检验动态身份连续性。’面对追问时，我会展开自己亲自完成的部分，并说明其他成员的职责。",
         "请将【真实模块】替换为你的准确分工。二作身份下，清楚边界比夸大贡献更重要。"),
        ("如果独立延续 EyeStyleMimic 的研究，你下一步会做什么？",
         "我会从三条线推进。第一，扩大参考图类型、用户群体和机器人呈现形式，验证结果是否能跨风格和人群复现。第二，引入专家设计参数或更强的生成基线，进一步分析 21 项参数中哪些特征最影响身份连续性。第三，把系统放进真实交互场景，研究身份连续性是否会影响用户的信任、陪伴感或互动投入。同时，论文中的校正器尚未证明能提升视觉质量，后续需要用独立数据和用户或专家评价检验哪些参数修正真正有效。",
         "后续工作必须与论文已报告的局限对应，避免提出与现有结果无关的泛泛方向。")
    ]
    for i, (q, a, t) in enumerate(research_paper_questions, 1):
        add_question(doc, f"研究{i}", q, a, t)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para_format(footer, before=4, after=0, line=1.0)
    r = footer.add_run("科研面试模拟问答  第 ")
    set_run_font(r, size=9, color=(100, 100, 100))
    add_page_field(footer)
    r = footer.add_run(" 页")
    set_run_font(r, size=9, color=(100, 100, 100))

    doc.core_properties.title = "科研面试模拟问答"
    doc.core_properties.author = "熊晨涵"
    doc.save(OUTPUT)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    main()
